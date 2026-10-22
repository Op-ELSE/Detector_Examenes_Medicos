import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Dynamically generate columns from fields.json
from pathlib import Path
import json

# Path to fields.json located in the project root
FIELDS_PATH = Path(__file__).parent / 'fields.json'
if FIELDS_PATH.exists():
    with FIELDS_PATH.open(encoding='utf-8') as f:
        FIELD_NAMES = json.load(f)
else:
    FIELD_NAMES = []

import re

def _normalize_date(val):
    if not val:
        return val
    s = str(val).strip()
    m1 = re.search(r'\b(19\d\d|20\d\d)[-/.](\d{1,2})[-/.](\d{1,2})\b', s)
    if m1:
        return f'{int(m1.group(2)):02d}/{int(m1.group(3)):02d}/{m1.group(1)}'
    m2 = re.search(r'\b(\d{1,2})[-/.](\d{1,2})[-/.](19\d\d|20\d\d)\b', s)
    if m2:
        return f'{int(m2.group(2)):02d}/{int(m2.group(1)):02d}/{m2.group(3)}'
    return s

def _normalize_aptitud(val):
    if not val:
        return val
    s = str(val).strip().upper()
    if 'NO APTO' in s:
        return 'NO APTO'
    if 'RESTRICCI' in s or 'OBSERV' in s:
        return 'APTO CON RESTRICCIONES'
    if 'APTO' in s:
        return 'APTO'
    return s

def _clean_numeric_value(val):
    if val is None:
        return val
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if not s:
        return s
    # Ignorar fechas
    if re.search(r'\b\d{1,2}/\d{1,2}/\d{4}\b', s) or re.search(r'\b\d{4}-\d{2}-\d{2}\b', s):
        return s
    if not re.search(r'\d', s):
        return s
    m = re.match(r'^\s*([><=~±]?\s*\d+(?:[.,]\d+)?(?:\s*/\s*\d+(?:[.,]\d+)?)?)\s*(?:[a-zA-Z%µ/°].*)?$', s)
    if m:
        return m.group(1).strip()
    cleaned = re.sub(r'(?i)\s*(g/dl|mg/dl|mmol/l|ui/l|u/l|u/ml|mmhg|cm|kg|m2|m|l|ml|%|porcentaje|x10\^?\d*|cel/campo|p/c|leuc/campo|fl|pg|g/l|seg|segundos|bpm|lpm|rpm)\s*$', '', s).strip()
    return cleaned if cleaned else s

def _append_comment_if_extra(std_choice, original_text):
    if not original_text:
        return std_choice
    orig_clean = str(original_text).strip()
    if ' - ' in orig_clean:
        partes = orig_clean.split(' - ', 1)
        det = partes[1].strip()
        if det:
            return f"{std_choice} - {det}"
        return std_choice

    orig_upper = orig_clean.upper()
    trivial_normals = [
        'NORMAL', 'SIN ALTERACIONES', 'S/A', 'SIN ALTERACION', 'CONSERVADO', 'CONSERVADA',
        'APARENTEMENTE NORMAL', 'ADECUADO', 'ADECUADA', 'BUENO', 'BUENA', 'VENTILAN BIEN',
        'CAMPOS PULMONARES LIMPIOS Y NORMALES', 'CAMPOS LIMPIOS', 'SIN LESIONES',
        'NO SE OBSERVAN ALTERACIONES', 'NEGATIVO', '0', '0 MG/DL', 'NO CONTIENE'
    ]
    if any(orig_upper == tn or orig_upper == f"NORMAL {tn}" for tn in trivial_normals):
        return std_choice

    cleaned_detail = re.sub(
        r'^\s*\b(normal|alterado|anormal|patologico|si|sí|se realizo|se realizó|no|positivo|negativo|conservado|conservada|limpido|límpido|turbio|sinusal|ritmo sinusal)\b\s*[-:,;.]*\s*',
        '',
        orig_clean,
        flags=re.IGNORECASE
    ).strip()

    if cleaned_detail and len(cleaned_detail) > 2 and cleaned_detail.upper() != std_choice.upper():
        cleaned_detail = cleaned_detail.strip(" '\"-:,;()[]+")
        if cleaned_detail:
            cleaned_detail = cleaned_detail[0].upper() + cleaned_detail[1:] if len(cleaned_detail) > 1 else cleaned_detail
            return f"{std_choice} - {cleaned_detail}"

    return std_choice

def _normalize_standard_values(field_id, val):
    if val is None:
        return val
    s = str(val).strip()
    if not s:
        return s
    
    key_lower = field_id.lower()
    if any(k in key_lower for k in ['recomendaci', 'observaci', 'conclusion', 'hallazgo', 'nombre', 'restricci', 'motivo', 'instrumento', 'centro_medico', 'comentario', 'firma', 'validacion']):
        return s

    s_upper = s.upper()

    # Aptitud
    if 'aptitud' in key_lower or key_lower in ['resultado_aptitud_final', 'alt_resultado_de_suficiencia']:
        return _normalize_aptitud(val)

    # Booleanos / Realizados
    if any(k in key_lower for k in ['realizada', 'realizado', 'validada', 'validado', 'declarado', 'cumple']):
        if any(w in s_upper for w in ['SI', 'SÍ', 'REALIZADO', 'REALIZADA', 'APLICADO', 'VALIDADO', 'CORRECTO', 'TRUE', '1']):
            return _append_comment_if_extra('SI', s)
        if any(w in s_upper for w in ['NO', 'NO REALIZADO', 'NO REALIZADA', 'FALSO', 'FALSE', '0']):
            return _append_comment_if_extra('NO', s)
        return s

    # Drogas
    if 'dro_' in key_lower or 'cocaina' in key_lower or 'cannabinoide' in key_lower:
        if any(w in s_upper for w in ['NO REACTIVO', 'NEGATIVO', 'NO DETECTADO']):
            return _append_comment_if_extra('NO REACTIVO', s)
        if any(w in s_upper for w in ['REACTIVO', 'POSITIVO', 'DETECTADO']):
            return _append_comment_if_extra('REACTIVO', s)
        if 'INDETERMINADO' in s_upper or 'DUDOSO' in s_upper:
            return _append_comment_if_extra('INDETERMINADO', s)

    # Grupo y factor
    if 'rh_grupo_sanguineo' in key_lower:
        m_abo = re.search(r'\b(A|B|AB|O)\b', s_upper)
        if m_abo:
            return m_abo.group(1)
    if 'rh_factor' in key_lower or 'factor_rh' in key_lower:
        if 'POS' in s_upper or '+' in s_upper:
            return 'POSITIVO (+)'
        if 'NEG' in s_upper or '-' in s_upper:
            return 'NEGATIVO (-)'

    # EKG
    if 'ekg_ritmo' in key_lower:
        if 'SINUSAL' in s_upper:
            return _append_comment_if_extra('SINUSAL', s)
        if 'ARRIT' in s_upper:
            return _append_comment_if_extra('ARRITMIA', s)
        if 'NO SINUSAL' in s_upper:
            return _append_comment_if_extra('NO SINUSAL', s)

    # Oftalmología
    if 'oft_resultado_de_ishihara' in key_lower or 'vision_de_colores' in key_lower:
        if any(w in s_upper for w in ['NORMAL', 'NEGATIVO', 'CONSERVAD', 'CONSERVA']):
            return _append_comment_if_extra('NORMAL', s)
        if any(w in s_upper for w in ['ALTERAD', 'DEFICIT', 'POSITIVO', 'DALTONISMO', 'ANORMAL']):
            return _append_comment_if_extra('ALTERADO', s)

    if 'oft_uso_de_lentes' in key_lower:
        if any(w in s_upper for w in ['SI', 'SÍ', 'USA', 'CORRECTOR']):
            return _append_comment_if_extra('SI', s)
        if any(w in s_upper for w in ['NO', 'NO USA']):
            return _append_comment_if_extra('NO', s)

    if 'oft_estereopsis' in key_lower:
        if any(w in s_upper for w in ['NORMAL', 'CONSERVAD', 'PRESENTE', 'POSITIV']):
            return _append_comment_if_extra('CONSERVADA', s)
        if any(w in s_upper for w in ['ALTERAD', 'ANORMAL', 'DISMINUID', 'NEGATIV']):
            return _append_comment_if_extra('ALTERADA', s)

    # Orina
    if 'ori_aspecto' in key_lower:
        if 'LIMPID' in s_upper or 'CLARO' in s_upper or 'TRANSPARENTE' in s_upper:
            return _append_comment_if_extra('LIMPIDO', s)
        if 'LIGERAMENTE TURBIO' in s_upper or 'SEMILIMPIDO' in s_upper:
            return _append_comment_if_extra('LIGERAMENTE TURBIO', s)
        if 'TURBIO' in s_upper:
            return _append_comment_if_extra('TURBIO', s)

    if 'ori_color' in key_lower:
        if 'AMARILLO' in s_upper or 'AMBAR' in s_upper:
            return _append_comment_if_extra('AMARILLO', s)
        if 'ROJIZO' in s_upper or 'HEMAT' in s_upper:
            return _append_comment_if_extra('ROJIZO', s)

    if any(k in key_lower for k in ['ori_glucosa', 'ori_hemoglobina', 'ori_proteinas', 'ori_sangre', 'ori_cetonas', 'ori_bilirrubina', 'ori_nitritos']):
        if any(w in s_upper for w in ['NEGATIVO', 'NO CONTIENE', 'AUSENTE', 'NO OBSERVA', '0', 'NORMAL']):
            return _append_comment_if_extra('NEGATIVO', s)
        if any(w in s_upper for w in ['POSITIVO', 'TRAZAS', 'PRESENTE', '+']):
            return _append_comment_if_extra('POSITIVO', s)

    # Órganos / Sistemas
    if any(key_lower.startswith(p) for p in ['cli_', 'rx_', 'mus_', 'der_', 'alt_']):
        if any(w in s_upper for w in ['NORMAL', 'CONSERVAD', 'SIN ALTERACI', 'S/A', 'NEGATIV', 'APARENTEMENTE NORMAL', 'ADECUAD', 'BUENO', 'VENTILAN BIEN']):
            return _append_comment_if_extra('NORMAL', s)
        return _append_comment_if_extra('ALTERADO', s)

    # Resultados generales
    if 'resultado' in key_lower or 'indicador' in key_lower:
        if any(w in s_upper for w in ['NORMAL', 'DENTRO DE LIMITES', 'DENTRO DE RANGO', 'NEGATIVO', 'SIN ALTERACION', 'CONSERVADO', 'ADECUADO', 'OPTIMO', 'NO PATOLOGICO']):
            return _append_comment_if_extra('NORMAL', s)
        if any(w in s_upper for w in ['ALTERADO', 'FUERA DE RANGO', 'PATOLOGICO', 'ANORMAL', 'OBSERVADO', 'POSITIVO']):
            return _append_comment_if_extra('ALTERADO', s)
        if any(w in s_upper for w in ['NO APTO']):
            return 'NO APTO'
        if any(w in s_upper for w in ['APTO CON RESTRIC', 'CON RESTRIC']):
            return 'APTO CON RESTRICCIONES'
        if any(w in s_upper for w in ['APTO']):
            return 'APTO'

    return s

def _make_columns():
    cols = []
    # Always include the PDF filename column
    cols.append(('archivo', 'Archivo PDF', 26))
    for field in FIELD_NAMES:
        cols.append((field, field, 22))
    return cols

COLUMNS = _make_columns()

def generate_excel(records):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Exámenes Médicos'
    ws.views.sheetView[0].showGridLines = True

    # Estilos de Encabezado
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='Segoe UI', size=10)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    wrap_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Colores para Aptitud Médica
    fill_apto = PatternFill(start_color='D1E7DD', end_color='D1E7DD', fill_type='solid')
    font_apto = Font(name='Segoe UI', size=10, bold=True, color='0F5132')

    fill_restriccion = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
    font_restriccion = Font(name='Segoe UI', size=10, bold=True, color='664D03')

    fill_no_apto = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
    font_no_apto = Font(name='Segoe UI', size=10, bold=True, color='842029')

    fill_observado = PatternFill(start_color='E2D9F3', end_color='E2D9F3', fill_type='solid')
    font_observado = Font(name='Segoe UI', size=10, bold=True, color='432874')

    fill_alt_row = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

    # Fila de Encabezados
    ws.row_dimensions[1].height = 32
    for col_idx, (field_id, header_text, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Filas de Datos
    for row_idx, rec in enumerate(records, 2):
        ws.row_dimensions[row_idx].height = 28
        is_alt = (row_idx % 2 == 1)

        for col_idx, (field_id, _, _) in enumerate(COLUMNS, 1):
            val = rec.get(field_id, '')
            if isinstance(val, list):
                val = '; '.join(str(x) for x in val)
            elif val is None:
                val = ''
            
            if 'fecha' in field_id.lower() and val:
                val = _normalize_date(val)
            else:
                val = _clean_numeric_value(val)
                val = _normalize_standard_values(field_id, val)
            
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val))
            cell.font = data_font
            cell.border = thin_border

            if 'fecha' in field_id.lower() or field_id in ['dni', 'edad', 'sexo', 'tipo_examen', 'fecha_examen', 'fecha_vencimiento', 'Fecha_Examen', 'Fecha_Vencimiento', 'DNI']:
                cell.alignment = center_align
            elif field_id in ['restricciones', 'diagnosticos', 'recomendaciones', 'Recomendacion_General', 'Restricciones_General']:
                cell.alignment = wrap_align
            else:
                cell.alignment = left_align

            if is_alt:
                cell.fill = fill_alt_row

            # Formato condicional para Aptitud
            if field_id in ['aptitud', 'Resultado_Aptitud_Final']:
                apt_upper = str(val).upper()
                cell.alignment = center_align
                if 'RESTRICCION' in apt_upper or 'OBSERVACIONES' in apt_upper:
                    cell.fill = fill_restriccion
                    cell.font = font_restriccion
                elif 'NO APTO' in apt_upper:
                    cell.fill = fill_no_apto
                    cell.font = font_no_apto
                elif 'OBSERVADO' in apt_upper or 'PENDIENTE' in apt_upper:
                    cell.fill = fill_observado
                    cell.font = font_observado
                elif 'APTO' in apt_upper:
                    cell.fill = fill_apto
                    cell.font = font_apto

    # Auto-filtro
    num_rows = max(len(records) + 1, 1)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{num_rows}'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer